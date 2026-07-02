"""Construye documentacion/presentacion_is_it_ai.pptx — tema claro, dinámico, 8 slides.

Estructura fija:
  1 Presentación · 2 Descripción del proyecto · 3 Separación en 3 categorías ·
  4 Datos usados · 5 PyTorch · 6 TensorFlow · 7 Estado actual general ·
  8 Próximos pasos

Diseño: orbes de gradiente difuminados de fondo, paneles con tinte translúcido,
números fantasma por sección, tarjetas con cabecera de color, cierre en azul
noche. Texto mínimo en pantalla; el detalle vive en notas_presentacion.md.

Animaciones automáticas (fade secuencial sin clicks) + transición fade rápida.

Corre con el venv de herramientas (tiene python-pptx):
  herramientas/.venv/Scripts/python.exe documentacion/deck_assets/build_pptx_light.py
"""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parents[1]
SALIDA = RAIZ / "documentacion" / "presentacion_is_it_ai.pptx"

BG = RGBColor(0xF7, 0xF8, 0xFC)
NOCHE = RGBColor(0x10, 0x15, 0x27)   # azul noche (cierre / tarjetas de contraste)
FG = RGBColor(0x1B, 0x21, 0x30)
MUT = RGBColor(0x5C, 0x65, 0x77)
VIO = RGBColor(0x6D, 0x4A, 0xFF)
TEA = RGBColor(0x0E, 0xA8, 0xA0)
TF_C = RGBColor(0xE8, 0x73, 0x1A)
PT_C = RGBColor(0xD9, 0x4A, 0x26)
ROJO = RGBColor(0xD9, 0x30, 0x25)
VERDE = RGBColor(0x16, 0xA3, 0x4A)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
BORDE = RGBColor(0xE2, 0xE8, 0xF0)
GHOST = RGBColor(0xE4, 0xE8, 0xF4)

FONT_H = "Segoe UI"
FONT_B = "Segoe UI"
ANCHO, ALTO = Inches(13.333), Inches(7.5)
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"


# ---------------------------------------------------------------- helpers ---
def fondo(slide, color=BG):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _set_alpha(shape, pct):
    """Transparencia del relleno sólido (pct = opacidad 0-100)."""
    srgb = shape.fill._xPr.find(".//" + qn("a:srgbClr"))
    a = srgb.makeelement(qn("a:alpha"), {"val": str(int(pct * 1000))})
    srgb.append(a)


def orb(slide, cual, x, y, size):
    """Orbe difuminado decorativo (violeta o teal)."""
    p = AQUI / f"orb_{cual}.png"
    pic = slide.shapes.add_picture(str(p), x, y, width=size, height=size)
    pic.line.fill.background()
    return pic


def panel(slide, x, y, w, h, color, alpha=100, borde=None, radio=True):
    shp = slide.shapes.add_shape(5 if radio else 1, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if alpha < 100:
        _set_alpha(shp, alpha)
    if borde:
        shp.line.color.rgb = borde
        shp.line.width = Pt(1.25)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def barra_gradiente(slide, y=Inches(1.16), x=Inches(0.75), w=Inches(2.4), h=Pt(4.5)):
    shp = slide.shapes.add_shape(1, x, y, w, h)
    shp.line.fill.background()
    shp.fill.gradient()
    stops = shp.fill.gradient_stops
    stops[0].color.rgb = VIO
    stops[0].position = 0.0
    stops[1].color.rgb = TEA
    stops[1].position = 1.0
    try:
        shp.fill.gradient_angle = 0
    except Exception:
        pass
    shp.shadow.inherit = False
    return shp


def texto(slide, x, y, w, h, contenido, size=18, color=FG, bold=False,
          align=PP_ALIGN.LEFT, font=FONT_B, line_spacing=1.12, anchor=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    if anchor:
        tf.vertical_anchor = anchor
    items = contenido if isinstance(contenido, list) else [(contenido, {})]
    for i, (txt, ov) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ov.get("align", align)
        p.line_spacing = ov.get("line_spacing", line_spacing)
        p.space_after = Pt(ov.get("space_after", 5))
        r = p.add_run()
        r.text = txt
        f = r.font
        f.name = ov.get("font", font)
        f.size = Pt(ov.get("size", size))
        f.bold = ov.get("bold", bold)
        f.color.rgb = ov.get("color", color)
        if ov.get("italic"):
            f.italic = True
    return tb


def kicker_chip(slide, txt, color=TEA):
    w = Inches(0.32 + 0.088 * len(txt))
    shp = panel(slide, Inches(0.75), Inches(0.30), w, Inches(0.36), color, alpha=13)
    tb = texto(slide, Inches(0.75), Inches(0.295), w, Inches(0.36),
               txt.upper(), size=12, color=color, bold=True,
               align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return [shp, tb]


def ghost(slide, n, color=GHOST):
    return texto(slide, Inches(10.6), Inches(-0.25), Inches(2.6), Inches(1.9),
                 f"{n:02d}", size=110, color=color, bold=True, font=FONT_H,
                 align=PP_ALIGN.RIGHT)


def titulo(slide, txt, kicker=None, n=None, color_txt=FG, kcolor=TEA):
    piezas = []
    if n is not None:
        ghost(slide, n)
    if kicker:
        piezas += kicker_chip(slide, kicker, kcolor)
    piezas.append(texto(slide, Inches(0.72), Inches(0.72), Inches(10.5), Inches(0.72),
                        txt, size=30, color=color_txt, bold=True, font=FONT_H))
    piezas.append(barra_gradiente(slide, y=Inches(1.35)))
    return piezas


def pie(slide, n, total, claro=False):
    c = RGBColor(0x8B, 0x93, 0xA8) if claro else MUT
    texto(slide, Inches(0.75), Inches(7.08), Inches(6), Inches(0.32),
          "Is-it-AI · Frameworks de IA · UDD", size=10, color=c)
    texto(slide, Inches(12.1), Inches(7.08), Inches(0.9), Inches(0.32),
          f"{n:02d} / {total}", size=10, color=c, align=PP_ALIGN.RIGHT)


def imagen(slide, ruta, x, y, w=None, h=None, borde=False):
    pic = slide.shapes.add_picture(str(ruta), x, y, width=w, height=h)
    if borde:
        pic.line.color.rgb = BORDE
        pic.line.width = Pt(1)
    else:
        pic.line.fill.background()
    return pic


def stat(slide, x, y, w, num, label, color):
    """Tarjeta de estadística: número grande + etiqueta."""
    card = panel(slide, x, y, w, Inches(1.55), CARD, borde=BORDE)
    top = panel(slide, x, y, w, Inches(0.10), color, radio=False)
    tb = texto(slide, x, y + Inches(0.18), w, Inches(1.3), [
        (num, {"size": 30, "bold": True, "color": color, "align": PP_ALIGN.CENTER}),
        (label, {"size": 12.5, "color": MUT, "align": PP_ALIGN.CENTER}),
    ])
    return [card, top, tb]


def notas(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt


# ------------------------------------------------- transiciones/animación ---
def transicion_fade(slide, thru_black=False):
    el = slide._element
    tr = etree.SubElement(el, f"{{{NS_P}}}transition")
    tr.set("spd", "fast")
    fade = etree.SubElement(tr, f"{{{NS_P}}}fade")
    if thru_black:
        fade.set("thruBlk", "1")


def _timing_xml(shape_ids, dur=450, stagger=260):
    cid = [7]

    def nid():
        cid[0] += 1
        return str(cid[0])

    efectos = []
    for i, spid in enumerate(shape_ids):
        delay = 150 if i == 0 else stagger
        efectos.append(f"""
          <p:par>
            <p:cTn id="{nid()}" presetID="10" presetClass="entr" presetSubtype="0"
                   fill="hold" grpId="0" nodeType="afterEffect">
              <p:stCondLst><p:cond delay="{delay}"/></p:stCondLst>
              <p:childTnLst>
                <p:set>
                  <p:cBhvr>
                    <p:cTn id="{nid()}" dur="1" fill="hold">
                      <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                    </p:cTn>
                    <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
                    <p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>
                  </p:cBhvr>
                  <p:to><p:strVal val="visible"/></p:to>
                </p:set>
                <p:animEffect transition="in" filter="fade">
                  <p:cBhvr>
                    <p:cTn id="{nid()}" dur="{dur}"/>
                    <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
                  </p:cBhvr>
                </p:animEffect>
              </p:childTnLst>
            </p:cTn>
          </p:par>"""
        )

    xml = f"""<p:timing xmlns:p="{NS_P}"
      xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
  <p:tnLst>
    <p:par>
      <p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">
        <p:childTnLst>
          <p:seq concurrent="1" nextAc="seek">
            <p:cTn id="2" dur="indefinite" nodeType="mainSeq">
              <p:childTnLst>
                <p:par>
                  <p:cTn id="3" fill="hold">
                    <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                    <p:childTnLst>
                      <p:par>
                        <p:cTn id="4" fill="hold">
                          <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                          <p:childTnLst>{"".join(efectos)}
                          </p:childTnLst>
                        </p:cTn>
                      </p:par>
                    </p:childTnLst>
                  </p:cTn>
                </p:par>
              </p:childTnLst>
            </p:cTn>
            <p:prevCondLst>
              <p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>
            </p:prevCondLst>
            <p:nextCondLst>
              <p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>
            </p:nextCondLst>
          </p:seq>
        </p:childTnLst>
      </p:cTn>
    </p:par>
  </p:tnLst>
</p:timing>"""
    return etree.fromstring(xml.encode())


def animar(slide, shapes, dur=450, stagger=260):
    ids = [s.shape_id for s in shapes]
    slide._element.append(_timing_xml(ids, dur, stagger))


# ------------------------------------------------------------------ build ---
prs = Presentation()
prs.slide_width = ANCHO
prs.slide_height = ALTO
BLANK = prs.slide_layouts[6]
TOTAL = 8


def nueva(color=BG):
    s = prs.slides.add_slide(BLANK)
    fondo(s, color)
    return s


# 1 — Presentación (portada Stitch light)
s = nueva()
src = Image.open(AQUI / "stitch_portada_light.png")
w0, h0 = src.size
h_obj = int(w0 * 9 / 16)
top = (h0 - h_obj) // 2
src.crop((0, top, w0, top + h_obj)).save(AQUI / "_portada_light_169.png")
imagen(s, AQUI / "_portada_light_169.png", 0, 0, w=ANCHO)
texto(s, Inches(0.75), Inches(6.98), Inches(8), Inches(0.4),
      "Frameworks de IA · UDD · Avance 1", size=12, color=MUT)
transicion_fade(s, thru_black=True)
notas(s, "Ver notas_presentacion.md — slide 1.")

# 2 — Descripción del proyecto
s = nueva()
orb(s, "violet", Inches(-2.3), Inches(4.2), Inches(5.5))
orb(s, "teal", Inches(10.8), Inches(-2.2), Inches(5.0))
titulo(s, "Un clasificador móvil de huella de IA", kicker="Descripción del proyecto", n=2)
b = texto(s, Inches(0.85), Inches(2.15), Inches(6.6), Inches(3.9), [
    ("Foto de una diapositiva → ¿cuánta huella de IA?", {"size": 22, "bold": True, "space_after": 16}),
    ("Medimos artefactos visuales, no procedencia.", {"size": 18, "space_after": 14}),
    ("CNN simple primero (v1 · v2) para diagnosticar;", {"size": 18, "color": MUT}),
    ("MobileNetV3 + export móvil después (v3).", {"size": 18, "color": MUT}),
])
t = panel(s, Inches(8.05), Inches(2.0), Inches(4.55), Inches(3.7), NOCHE)
tx = texto(s, Inches(8.45), Inches(2.45), Inches(3.8), Inches(2.9), [
    ("LA IDEA CLAVE", {"size": 12, "bold": True, "color": TEA}),
    ("Una foto contiene el resultado, no el proceso.", {"size": 24, "bold": True, "color": CARD}),
    ("El label dice lo que el modelo puede aprender.", {"size": 14, "color": RGBColor(0x9A, 0xA3, 0xB8)}),
])
ch1 = panel(s, Inches(0.85), Inches(5.55), Inches(2.3), Inches(0.5), TF_C, alpha=13)
ct1 = texto(s, Inches(0.85), Inches(5.55), Inches(2.3), Inches(0.5), "TensorFlow",
            size=14, color=TF_C, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
ch2 = panel(s, Inches(3.35), Inches(5.55), Inches(2.3), Inches(0.5), PT_C, alpha=13)
ct2 = texto(s, Inches(3.35), Inches(5.55), Inches(2.3), Inches(0.5), "PyTorch",
            size=14, color=PT_C, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
pie(s, 2, TOTAL)
transicion_fade(s)
animar(s, [b, t, tx, ch1, ct1, ch2, ct2])
notas(s, "Ver notas_presentacion.md — slide 2.")

# 3 — Separación en 3 categorías
s = nueva()
orb(s, "teal", Inches(-2.0), Inches(-2.3), Inches(5.2))
orb(s, "violet", Inches(10.6), Inches(4.6), Inches(5.2))
titulo(s, "Tres categorías — un eje ordinal", kicker="Separación en 3 categorías", n=3)
g = imagen(s, AQUI / "light_gradiente_clases.png", Inches(1.1), Inches(2.15), w=Inches(11.1))
f = panel(s, Inches(2.35), Inches(5.55), Inches(8.6), Inches(0.72), VIO, alpha=10)
ft = texto(s, Inches(2.35), Inches(5.55), Inches(8.6), Inches(0.72),
           "No es IA sí/no: es cuánto se nota — el softmax entrega confianza por nivel.",
           size=16, color=FG, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
pie(s, 3, TOTAL)
transicion_fade(s)
animar(s, [g, f, ft])
notas(s, "Ver notas_presentacion.md — slide 3.")

# 4 — Datos usados
s = nueva()
orb(s, "violet", Inches(11.0), Inches(-2.0), Inches(4.8))
orb(s, "teal", Inches(-2.2), Inches(4.8), Inches(5.0))
titulo(s, "Datos etiquetados por procedencia", kicker="Datos usados", n=4)
g = imagen(s, AQUI / "light_pipeline.png", Inches(0.7), Inches(1.95), w=Inches(11.9))
st1 = stat(s, Inches(1.35), Inches(5.0), Inches(3.2), "300", "imágenes reales", VIO)
st2 = stat(s, Inches(5.05), Inches(5.0), Inches(3.2), "~100", "por clase, balanceado", TEA)
st3 = stat(s, Inches(8.75), Inches(5.0), Inches(3.2), "80 / 20", "split train · val (seed fija)", VIO)
pie(s, 4, TOTAL)
transicion_fade(s)
animar(s, [g] + st1 + st2 + st3, stagger=200)
notas(s, "Ver notas_presentacion.md — slide 4.")


# 5/6 — Frameworks
def slide_framework(nombre, color, curvas, matriz, v1_stats, v2_stats, punchline, n):
    s = nueva()
    orb(s, "violet" if nombre == "PyTorch" else "teal", Inches(10.9), Inches(4.7), Inches(4.6))
    banda = panel(s, 0, 0, ANCHO, Inches(1.62), color, alpha=9, radio=False)
    titulo(s, f"{nombre} — resultados v1 → v2", kicker=f"Frameworks · {nombre}", n=n, kcolor=color)
    lt = texto(s, Inches(0.7), Inches(1.78), Inches(5), Inches(0.3),
               [("Curvas v2", {"size": 12.5, "bold": True, "color": color})])
    g1 = imagen(s, curvas, Inches(0.7), Inches(2.1), w=Inches(7.5), borde=True)
    lt2 = texto(s, Inches(0.7), Inches(4.6), Inches(5), Inches(0.3),
                [("Matriz de confusión v2 (val)", {"size": 12.5, "bold": True, "color": color})])
    g2 = imagen(s, matriz, Inches(0.7), Inches(4.92), h=Inches(2.0), borde=True)
    pl = panel(s, Inches(3.35), Inches(5.35), Inches(4.85), Inches(1.1), color, alpha=11)
    plt_ = texto(s, Inches(3.55), Inches(5.35), Inches(4.5), Inches(1.1), punchline,
                 size=14.5, color=FG, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    # Tarjeta v1 (overfit) y v2 (regularizado)
    c1 = panel(s, Inches(8.6), Inches(2.1), Inches(4.05), Inches(2.15), CARD, borde=BORDE)
    h1 = panel(s, Inches(8.6), Inches(2.1), Inches(4.05), Inches(0.42), ROJO, alpha=14)
    x1 = texto(s, Inches(8.85), Inches(2.12), Inches(3.6), Inches(2.0), [
        ("v1 · sin augmentation", {"size": 13, "bold": True, "color": ROJO}),
        (v1_stats[0], {"size": 21, "bold": True, "space_after": 0}),
        (v1_stats[1], {"size": 13, "color": MUT}),
    ])
    c2 = panel(s, Inches(8.6), Inches(4.5), Inches(4.05), Inches(2.15), CARD, borde=BORDE)
    h2 = panel(s, Inches(8.6), Inches(4.5), Inches(4.05), Inches(0.42), VERDE, alpha=14)
    x2 = texto(s, Inches(8.85), Inches(4.52), Inches(3.6), Inches(2.0), [
        ("v2 · con augmentation", {"size": 13, "bold": True, "color": VERDE}),
        (v2_stats[0], {"size": 21, "bold": True, "space_after": 0}),
        (v2_stats[1], {"size": 13, "color": MUT}),
    ])
    pie(s, n, TOTAL)
    transicion_fade(s)
    animar(s, [lt, g1, c1, h1, x1, c2, h2, x2, lt2, g2, pl, plt_], stagger=220)
    notas(s, f"Ver notas_presentacion.md — slide {n}.")


slide_framework(
    "PyTorch", PT_C,
    RAIZ / "PyTorch" / "v2" / "Figure_1.png",
    RAIZ / "PyTorch" / "v2" / "Figure_2_matriz.png",
    ("0.996 → 0.833 acc", "loss 0.029 vs 0.406: memoriza"),
    ("0.708 → 0.717 acc", "loss 0.669 vs 0.707: sin brecha"),
    [("Augmentation solo en el loader de train. Error entre vecinos; 1 caso extremo.", {})],
    5,
)

slide_framework(
    "TensorFlow", TF_C,
    RAIZ / "TensorFlow" / "v2" / "Figure_1.png",
    RAIZ / "TensorFlow" / "v2" / "Figure_2_matriz.png",
    ("1.000 → 0.967 acc", "loss 0.010 vs 0.103: brecha ~10×"),
    ("0.771 → 0.750 acc", "loss 0.529 vs 0.550: sin brecha"),
    [("Augmentation como capas del modelo. 3 casos extremos 2→0: re-evaluar al converger.", {})],
    6,
)

# 7 — Estado actual general
s = nueva()
orb(s, "teal", Inches(10.9), Inches(-2.1), Inches(4.8))
orb(s, "violet", Inches(-2.2), Inches(4.6), Inches(5.0))
titulo(s, "Dónde estamos", kicker="Estado actual general", n=7)


def columna(x, color, encabezado, items):
    card = panel(s, x, Inches(2.0), Inches(3.9), Inches(4.5), CARD, borde=BORDE)
    head = panel(s, x, Inches(2.0), Inches(3.9), Inches(0.62), color, radio=False)
    ht = texto(s, x, Inches(2.0), Inches(3.9), Inches(0.62), encabezado, size=16,
               bold=True, color=CARD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    cuerpo = [(f"· {t}", {"size": 14.5, "space_after": 10}) for t in items]
    bt = texto(s, x + Inches(0.3), Inches(2.85), Inches(3.3), Inches(3.5), cuerpo)
    return [card, head, ht, bt]


c1 = columna(Inches(0.75), VERDE, "Logrado", [
    "Pipeline de datos completo",
    "Dataset 300 imgs balanceado",
    "v1 y v2 en ambos frameworks",
    "Matriz de confusión operando",
])
c2 = columna(Inches(4.85), TEA, "Diagnóstico", [
    "v1 sobreajusta (loss ~10×)",
    "v2 lo elimina (augmentation)",
    "Quedó sub-entrenado",
    "2 frameworks → misma lectura",
])
c3 = columna(Inches(8.95), ROJO, "Pendiente", [
    "Protocolo frontera 1↔2",
    "Domain gap (fotos reales)",
    "Confusión 0↔2 (¿ruido?)",
    "MobileNetV3 + móvil (v3)",
])
pie(s, 7, TOTAL)
transicion_fade(s)
animar(s, c1 + c2 + c3, stagger=180)
notas(s, "Ver notas_presentacion.md — slide 7.")

# 8 — Próximos pasos (azul noche)
s = nueva(NOCHE)
orb(s, "violet", Inches(-2.5), Inches(-2.5), Inches(6.5))
orb(s, "teal", Inches(10.3), Inches(3.9), Inches(6.0))
kicker_chip(s, "Próximos pasos", TEA)
texto(s, Inches(0.72), Inches(0.72), Inches(10.5), Inches(0.72),
      "Qué sigue", size=30, color=CARD, bold=True, font=FONT_H)
barra_gradiente(s, y=Inches(1.35))
g = imagen(s, AQUI / "light_roadmap.png", Inches(0.7), Inches(2.0), w=Inches(11.9))
m = texto(s, Inches(0.9), Inches(5.05), Inches(11.5), Inches(0.7),
          "Leer el estado del modelo y responder — no amontonar técnicas.",
          size=17, bold=True, color=RGBColor(0xC9, 0xD0, 0xE2), align=PP_ALIGN.CENTER)
q = texto(s, Inches(0.9), Inches(5.75), Inches(11.5), Inches(1.1),
          "¿Preguntas?", size=44, bold=True, color=TEA, align=PP_ALIGN.CENTER, font=FONT_H)
pie(s, 8, TOTAL, claro=True)
transicion_fade(s, thru_black=True)
animar(s, [g, m, q])
notas(s, "Ver notas_presentacion.md — slide 8.")

prs.save(SALIDA)
print(f"OK -> {SALIDA}")
